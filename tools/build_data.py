from __future__ import annotations

import json
import math
import re
import shutil
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORK_ROOT = ROOT.parent
SOURCE = WORK_ROOT / "v5-recovery-source"
OUT_DATA = ROOT / "data"

EXCEL = SOURCE / "PR – V1.xlsx"
TEXT_JSON = SOURCE / "PRX_Kurztexte_PR_POI_V1.json"
GPX_ZIP = SOURCE / "Gpx (1).zip"
KML_ZIP = SOURCE / "KML (1).zip"


def slug_pr(value: str) -> str:
    raw = str(value or "").upper().replace("_", " ").replace("-", " ")
    match = re.search(r"(?:PS\s*)?PR\s*([0-9]+(?:\.[0-9]+)?)", raw)
    if not match:
        match = re.search(r"([0-9]+(?:\.[0-9]+)?)", raw)
    if not match:
        return re.sub(r"[^A-Z0-9]+", "-", raw).strip("-").lower()
    prefix = "ps-pr" if "PS" in raw else "pr"
    return f"{prefix}-{match.group(1).replace('.', '-')}"


def display_pr(value: str) -> str:
    raw = str(value or "").replace("_", " ").strip()
    raw = re.sub(r"\s+", " ", raw)
    raw = raw.replace("PR ", "PR")
    raw = raw.replace("PR", "PR ")
    raw = re.sub(r"\s+", " ", raw).strip()
    return raw


def number_pr(value: str) -> str:
    raw = display_pr(value)
    return raw.replace("PR ", "").strip()


def to_float(value):
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", "."))
    except ValueError:
        return None


def to_int(value):
    num = to_float(value)
    if num is None or not math.isfinite(num):
        return None
    return int(round(num))


def read_text_bundle():
    if not TEXT_JSON.exists():
        return {}, []
    data = json.loads(TEXT_JSON.read_text(encoding="utf-8-sig"))
    prs = {}
    for item in data.get("prs", []):
        prs[slug_pr(item.get("id") or item.get("routeNo"))] = item
    return prs, data.get("pois", [])


def hyperlink(cell):
    return cell.hyperlink.target if cell.hyperlink else None


def read_excel_prs():
    wb = load_workbook(EXCEL, read_only=False, data_only=True)
    ws = wb.active
    headers = [str(c.value or "").strip() for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2):
        values = {headers[idx]: cell.value for idx, cell in enumerate(row) if idx < len(headers)}
        if not values.get("Nummer") or not values.get("Name"):
            continue
        route_id = display_pr(values.get("Nummer"))
        rows.append(
            {
                "id": slug_pr(route_id),
                "displayId": route_id,
                "number": number_pr(route_id),
                "name": str(values.get("Name") or "").strip(),
                "region": str(values.get("Region") or "").strip(),
                "lat": to_float(values.get("Latitude")),
                "lon": to_float(values.get("Longitude")),
                "distanceKm": to_float(values.get("Distanz")),
                "duration": str(values.get("Dauer") or "").strip(),
                "difficulty": str(values.get("Level") or "").strip(),
                "driveKm": to_float(values.get("Anfahrt km")),
                "driveMin": to_int(values.get("Anfaht Dauer")),
                "parking": str(values.get("Parkplatz") or "").strip(),
                "fee": values.get("Gebühr"),
                "elevationHigh": to_int(values.get("Höchster\nPunkt")),
                "elevationLow": to_int(values.get("Tiefster\nPunkt")),
                "elevationGain": to_int(values.get("Höhen-meter")),
                "links": {
                    "visitMadeira": hyperlink(row[0]),
                    "startGoogleMaps": hyperlink(row[7]) if len(row) > 7 else None,
                    "driveGoogleMaps": hyperlink(row[8]) if len(row) > 8 else None,
                    "schmalePfade": hyperlink(row[20]) if len(row) > 20 else None,
                },
            }
        )
    return rows


def read_zip_xml(zip_path: Path, ext: str):
    items = []
    with zipfile.ZipFile(zip_path) as zf:
        for info in zf.infolist():
            name = info.filename
            if name.startswith("__MACOSX/") or Path(name).name.startswith("._"):
                continue
            if not name.lower().endswith(ext):
                continue
            content = zf.read(info)
            items.append((Path(name).name, content))
    return items


def parse_gpx_points(content: bytes):
    root = ET.fromstring(content)
    points = []
    for elem in root.iter():
        if not elem.tag.endswith("trkpt"):
            continue
        lat = to_float(elem.attrib.get("lat"))
        lon = to_float(elem.attrib.get("lon"))
        ele = None
        for child in elem:
            if child.tag.endswith("ele"):
                ele = to_float(child.text)
                break
        if lat is not None and lon is not None:
            points.append([lat, lon, ele])
    return points


def parse_kml_points(content: bytes):
    root = ET.fromstring(content)
    points = []
    for elem in root.iter():
        if not elem.tag.endswith("coordinates") or not elem.text:
            continue
        for chunk in elem.text.split():
            parts = chunk.split(",")
            if len(parts) < 2:
                continue
            lon = to_float(parts[0])
            lat = to_float(parts[1])
            if lat is not None and lon is not None:
                points.append([lat, lon])
    if points:
        return points
    for elem in root.iter():
        if not elem.tag.endswith("coord") or not elem.text:
            continue
        parts = elem.text.split()
        if len(parts) < 2:
            continue
        lon = to_float(parts[0])
        lat = to_float(parts[1])
        if lat is not None and lon is not None:
            points.append([lat, lon])
    return points


def route_from_filename(name: str) -> str:
    return slug_pr(name)


def path_distance_km(points):
    total = 0.0
    for a, b in zip(points, points[1:]):
        total += haversine(a[0], a[1], b[0], b[1])
    return round(total, 2)


def haversine(lat1, lon1, lat2, lon2):
    radius = 6371.0088
    p1 = math.radians(lat1)
    p2 = math.radians(lat2)
    dp = math.radians(lat2 - lat1)
    dl = math.radians(lon2 - lon1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * radius * math.atan2(math.sqrt(a), math.sqrt(1 - a))


def write_geo_files():
    tracks = {}
    routes = {}
    for filename, content in read_zip_xml(GPX_ZIP, ".gpx"):
        rid = route_from_filename(filename)
        points = parse_gpx_points(content)
        if not points:
            continue
        out = {"id": rid, "sourceFile": filename, "type": "gpx", "points": points, "distanceKm": path_distance_km(points)}
        (OUT_DATA / "tracks" / f"{rid}.json").write_text(json.dumps(out, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        tracks[rid] = {"file": f"data/tracks/{rid}.json", "sourceFile": filename, "points": len(points), "distanceKm": out["distanceKm"]}
    for filename, content in read_zip_xml(KML_ZIP, ".kml"):
        rid = route_from_filename(filename)
        points = parse_kml_points(content)
        if not points:
            continue
        out = {"id": rid, "sourceFile": filename, "type": "kml", "points": points, "distanceKm": path_distance_km(points)}
        (OUT_DATA / "routes" / f"{rid}.json").write_text(json.dumps(out, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        routes[rid] = {"file": f"data/routes/{rid}.json", "sourceFile": filename, "points": len(points), "distanceKm": out["distanceKm"]}
    return tracks, routes


def main():
    shutil.rmtree(OUT_DATA / "tracks", ignore_errors=True)
    shutil.rmtree(OUT_DATA / "routes", ignore_errors=True)
    (OUT_DATA / "tracks").mkdir(parents=True, exist_ok=True)
    (OUT_DATA / "routes").mkdir(parents=True, exist_ok=True)
    text_prs, pois = read_text_bundle()
    prs = read_excel_prs()
    tracks, routes = write_geo_files()
    for pr in prs:
        extra = text_prs.get(pr["id"], {})
        pr["shortText"] = extra.get("short_150") or ""
        pr["detailText"] = extra.get("detail_280") or ""
        pr["featureTags"] = [x.strip() for x in str(extra.get("feature_tags") or "").split(",") if x.strip()]
        pr["status"] = extra.get("status_2026_06_07") or "Check"
        pr["track"] = tracks.get(pr["id"])
        pr["route"] = routes.get(pr["id"])
    data = {
        "meta": {
            "version": "V5.0.0 Recovery Core",
            "generatedFrom": ["PR – V1.xlsx", "PRX_Kurztexte_PR_POI_V1.json", "Gpx (1).zip", "KML (1).zip"],
            "counts": {"prs": len(prs), "tracks": len(tracks), "routes": len(routes), "pois": len(pois)},
        },
        "prs": prs,
    }
    (OUT_DATA / "prs.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT_DATA / "pois.json").write_text(json.dumps({"meta": {"version": "V5.0.0 Recovery Core"}, "pois": pois}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(data["meta"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
